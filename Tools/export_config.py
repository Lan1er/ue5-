# -*- coding: utf-8 -*-
"""《数值配置表.xlsx》 -> Content/Config/game_config.json
用法: python Tools/export_config.py
策划改完 Excel 后运行本脚本即可。校验失败会报错并中断导出。
"""
import json
import os
import sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "ConfigSource", "数值配置表.xlsx")
OUT_DIR = os.path.join(ROOT, "Content", "Config")

SCHEDULE_SHEETS = {"ScheduleWife": "Wife", "ScheduleFather": "Father"}

errors = []
warnings = []

def fail(msg):
    errors.append(msg)

def warn(msg):
    warnings.append(msg)

def to_minutes(s, ctx):
    """'HH:MM' -> 当天分钟数"""
    try:
        h, m = str(s).strip().split(":")
        v = int(h) * 60 + int(m)
        if not (0 <= v < 1440):
            fail(f"{ctx}: 时刻越界 '{s}'")
        return v
    except Exception:
        fail(f"{ctx}: 时刻格式错误 '{s}'（应为 HH:MM）")
        return 0

def num(v, ctx, lo=None, hi=None):
    if isinstance(v, bool):
        fail(f"{ctx}: 期望数字，得到布尔")
        return 0.0
    try:
        f = float(v)
    except Exception:
        fail(f"{ctx}: 期望数字，得到 '{v}'")
        return 0.0
    if lo is not None and f < lo:
        fail(f"{ctx}: {f} 小于下限 {lo}")
    if hi is not None and f > hi:
        fail(f"{ctx}: {f} 大于上限 {hi}")
    return f

def boolv(v, ctx):
    s = str(v).strip().upper()
    if s in ("TRUE", "1", "是"):
        return True
    if s in ("FALSE", "0", "否"):
        return False
    fail(f"{ctx}: 期望 TRUE/FALSE，得到 '{v}'")
    return False

def read_kv(ws, ctx):
    out = {}
    headers = [c.value for c in ws[1]]
    for i, col in enumerate(headers):
        if str(col).strip() == "Key":
            k_idx = i
        if str(col).strip() == "Value":
            v_idx = i
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[k_idx] is None:
            continue
        out[str(row[k_idx]).strip()] = row[v_idx]
    return out

def read_table(ws, ctx):
    headers = [str(c.value).strip() for c in ws[1]]
    rows = []
    for rn, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None for v in row):
            continue
        d = {}
        for h, v in zip(headers, row):
            d[h] = v
        d["_row"] = rn
        rows.append(d)
    return headers, rows

def main():
    if not os.path.exists(XLSX):
        sys.exit(f"找不到 {XLSX}")
    wb = load_workbook(XLSX, data_only=True)

    # ---- Global ----
    kv = read_kv(wb["Global"], "Global")
    gctx = "Global"
    global_cfg = {
        "DayStartMinutes":          to_minutes(kv.get("DayStartMinutes"), gctx + ".DayStartMinutes"),
        "NightMinutes":             to_minutes(kv.get("NightMinutes"), gctx + ".NightMinutes"),
        "GameMinutesPerRealSecond": num(kv.get("GameMinutesPerRealSecond"), gctx + ".GameMinutesPerRealSecond", lo=0.01),
        "CaughtResetsLoop":         boolv(kv.get("CaughtResetsLoop"), gctx + ".CaughtResetsLoop"),
        "NightEndsLoop":            boolv(kv.get("NightEndsLoop"), gctx + ".NightEndsLoop"),
        "MaxSuspicion":             num(kv.get("MaxSuspicion"), gctx + ".MaxSuspicion", lo=1),
        "MaxMood":                  num(kv.get("MaxMood"), gctx + ".MaxMood", lo=1),
    }
    if global_cfg["DayStartMinutes"] >= global_cfg["NightMinutes"]:
        fail("Global: 开始时刻必须早于入夜时刻")

    # ---- Suspicion ----
    kv = read_kv(wb["Suspicion"], "Suspicion")
    sctx = "Suspicion"
    sus = {
        "MaxValue":         num(kv.get("MaxValue"), sctx + ".MaxValue", lo=1),
        "DecayPerSecond":   num(kv.get("DecayPerSecond"), sctx + ".DecayPerSecond", lo=0),
        "DecayDelay":       num(kv.get("DecayDelay"), sctx + ".DecayDelay", lo=0),
        "ThresholdSearch":  num(kv.get("ThresholdSearch"), sctx + ".ThresholdSearch", lo=0),
        "ThresholdExposed": num(kv.get("ThresholdExposed"), sctx + ".ThresholdExposed", lo=0),
        "CaughtValue":      num(kv.get("CaughtValue"), sctx + ".CaughtValue", lo=1),
        "LoopResetDelta":   num(kv.get("LoopResetDelta"), sctx + ".LoopResetDelta"),
    }
    ths = ["ThresholdSearch", "ThresholdExposed", "CaughtValue"]
    for a, b in zip(ths, ths[1:]):
        if sus[a] >= sus[b]:
            fail(f"Suspicion: {a}({sus[a]}) 必须小于 {b}({sus[b]})")
    if sus["CaughtValue"] > sus["MaxValue"]:
        fail("Suspicion: CaughtValue 不能超过 MaxValue")

    headers, rows = read_table(wb["SuspicionGain"], "SuspicionGain")
    gains = []
    seen = set()
    for r in rows:
        t = str(r.get("TriggerType", "")).strip()
        if not t:
            fail(f"SuspicionGain 第{r['_row']}行: TriggerType 为空")
            continue
        if t in seen:
            fail(f"SuspicionGain 第{r['_row']}行: TriggerType '{t}' 重复")
        seen.add(t)
        g = num(r.get("Gain"), f"SuspicionGain.{t}.Gain")
        if g > 0 and not (5 <= g <= 20):
            warn(f"SuspicionGain.{t}.Gain 超出策划案建议区间 5~20")
        gains.append({"TriggerType": t, "Gain": g})
    sus["GainEntries"] = gains

    # ---- Mood + MoodStages ----
    kv = read_kv(wb["Mood"], "Mood")
    mctx = "Mood"
    mood = {
        "MaxValue":                       num(kv.get("MaxValue"), mctx + ".MaxValue", lo=1),
        "StressEventGain":                num(kv.get("StressEventGain"), mctx + ".StressEventGain", lo=0),
        "LegPainPerGameHour":             num(kv.get("LegPainPerGameHour"), mctx + ".LegPainPerGameHour", lo=0),
        "FoodRestore":                    num(kv.get("FoodRestore"), mctx + ".FoodRestore", lo=0),
        "LoopResetDelta":                 num(kv.get("LoopResetDelta"), mctx + ".LoopResetDelta"),
        "TrenchCoatMaxStage":             int(num(kv.get("TrenchCoatMaxStage"), mctx + ".TrenchCoatMaxStage", lo=1)),
        "Stage2PlusNoCoatSusGainMult":    num(kv.get("Stage2PlusNoCoatSusGainMult"), mctx + ".Stage2PlusNoCoatSusGainMult", lo=1),
        "Stage2PlusNoCoatSightRangeMult": num(kv.get("Stage2PlusNoCoatSightRangeMult"), mctx + ".Stage2PlusNoCoatSightRangeMult", lo=1),
    }
    # 2026-10-10：TrenchCoatBreakStage（风衣失效阶段）已删除。
    # 「风衣失效」规则在 2026-10-06-3 就被删掉了，这个字段和"必须小于"的校验都是历史残留。
    # 风衣现在只由 TrenchCoatMaxStage（可披上的最高体型阶段）决定，披上后不失效。

    headers, rows = read_table(wb["MoodStages"], "MoodStages")
    stages = []
    prev_min = -1
    for r in sorted(rows, key=lambda x: num(x.get("Stage"), "MoodStages.Stage")):
        st = int(num(r.get("Stage"), "MoodStages.Stage", lo=1, hi=9))
        mn = num(r.get("MinMood"), f"MoodStages.Stage{st}.MinMood", lo=0)
        if mn <= prev_min:
            fail(f"MoodStages: 阶段 {st} 的 MinMood({mn}) 必须大于前一阶段({prev_min})")
        prev_min = mn
        stages.append({
            "Stage": int(st),
            "MinMood": mn,
            "SizeMultiplier": num(r.get("SizeMultiplier"), f"MoodStages.Stage{st}.SizeMultiplier", lo=0.1),
            "Note": str(r.get("Note") or ""),
        })
    if stages and stages[0]["MinMood"] != 0:
        fail("MoodStages: 阶段 1 的 MinMood 必须为 0")
    mood["Stages"] = stages

    # ---- Periods（时段与探索机会） ----
    headers, rows = read_table(wb["Periods"], "Periods")
    periods = []
    prev_end = -1
    for r in rows:
        rn = r["_row"]
        start = to_minutes(r.get("Start"), f"Periods 第{rn}行.Start")
        end = to_minutes(r.get("End"), f"Periods 第{rn}行.End")
        if start >= end:
            fail(f"Periods 第{rn}行: Start 必须早于 End")
        if start < prev_end:
            fail(f"Periods 第{rn}行: 与上一时段重叠")
        prev_end = end
        pid = str(r.get("PeriodId") or "").strip()
        if not pid:
            fail(f"Periods 第{rn}行: PeriodId 为空")
        periods.append({
            "PeriodId": pid,
            "StartMinutes": start,
            "EndMinutes": end,
            "ExploreChances": int(num(r.get("ExploreChances"), f"Periods.{pid}.ExploreChances", lo=0)),
            "Note": str(r.get("Note") or ""),
        })
    if periods:
        if periods[0]["StartMinutes"] > global_cfg["DayStartMinutes"]:
            warn("Periods: 第一时段晚于一天开始时刻")
        if periods[-1]["EndMinutes"] > global_cfg["NightMinutes"]:
            warn("Periods: 最后时段超出入夜时刻")

    # ---- Items（道具表，按策划《道具(1).xlsx》录入） ----
    headers, rows = read_table(wb["Items"], "Items")
    items = []
    seen_items = set()
    valid_categories = {"Tool", "Food", "Medicine", "Story"}
    valid_storage = {"None", "Durable", "Perishable", "Unstorable"}
    for r in rows:
        rn = r["_row"]
        item_id = str(r.get("ItemId") or "").strip()
        if not item_id:
            fail(f"Items 第{rn}行: ItemId 为空")
            continue
        if item_id in seen_items:
            fail(f"Items 第{rn}行: ItemId '{item_id}' 重复")
        seen_items.add(item_id)
        if not (item_id.isascii() and item_id[0].isalpha()
                and all(ch.isalnum() or ch == "_" for ch in item_id)):
            fail(f"Items 第{rn}行: ItemId '{item_id}' 必须是英文标识符（PascalCase，字母开头）")
        display = str(r.get("DisplayName") or "").strip()
        if not display:
            fail(f"Items.{item_id}: DisplayName 为空")
        category = str(r.get("Category") or "").strip()
        if category not in valid_categories:
            fail(f"Items.{item_id}.Category: '{category}' 不在 {sorted(valid_categories)} 内")
        storage = str(r.get("StorageType") or "None").strip() or "None"
        if storage not in valid_storage:
            fail(f"Items.{item_id}.StorageType: '{storage}' 不在 {sorted(valid_storage)} 内")
        color = str(r.get("PlaceholderColor") or "").strip().lstrip("#")
        if color and not (len(color) == 6 and all(ch in "0123456789abcdefABCDEF" for ch in color)):
            fail(f"Items.{item_id}.PlaceholderColor: 需要 #RRGGBB（当前 '{color}'）")
            color = ""
        items.append({
            "ItemId": item_id,
            "DisplayName": display,
            "Category": category,
            "MaxStack": int(num(r.get("MaxStack"), f"Items.{item_id}.MaxStack", lo=1, hi=99)),
            "bCarryable": boolv(r.get("bCarryable"), f"Items.{item_id}.bCarryable"),
            "bConsumable": boolv(r.get("bConsumable"), f"Items.{item_id}.bConsumable"),
            "MoodDelta": num(r.get("MoodDelta"), f"Items.{item_id}.MoodDelta", lo=-100, hi=100),
            "SuspicionDelta": num(r.get("SuspicionDelta"), f"Items.{item_id}.SuspicionDelta", lo=-100, hi=100),
            "SpeedBonusPercent": num(r.get("SpeedBonusPercent"), f"Items.{item_id}.SpeedBonusPercent", lo=-100, hi=500),
            "SpiderDefense": num(r.get("SpiderDefense"), f"Items.{item_id}.SpiderDefense", lo=0, hi=10),
            "DoorBlockSeconds": num(r.get("DoorBlockSeconds"), f"Items.{item_id}.DoorBlockSeconds", lo=0, hi=600),
            "StorageType": storage,
            "PlaceholderColor": color.upper(),
            "AppearanceNote": str(r.get("AppearanceNote") or "").strip(),
            "SpawnLocation": str(r.get("SpawnLocation") or "").strip(),
            "Description": str(r.get("Description") or ""),
            # 剧情道具的后续剧情节点（策划对话表 B 组：读完回忆文本后播的人物台词）。
            # 留空 = 读完回忆文本直接拾取。指向的节点在 Story 解析完后统一交叉校验。
            "FollowUpStoryNode": str(r.get("FollowUpStoryNode") or "").strip(),
            "Note": str(r.get("Note") or ""),
        })
    if items and not any(i["MoodDelta"] < 0 for i in items if i["Category"] == "Food"):
        warn("Items: 没有任何「食物」带 MoodDelta<0（降低躁郁值），确认是否漏填")

    # ---- Schedules ----
    schedules = []
    for sheet_name, family_id in SCHEDULE_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue  # 动线 sheet 为可选：10-07 动线表定稿后加回 ScheduleWife/ScheduleFather
        headers, rows = read_table(wb[sheet_name], sheet_name)
        entries = []
        prev_end = -1
        for r in rows:
            rn = r["_row"]
            start = to_minutes(r.get("Start"), f"{sheet_name} 第{rn}行.Start")
            end = to_minutes(r.get("End"), f"{sheet_name} 第{rn}行.End")
            if start >= end:
                fail(f"{sheet_name} 第{rn}行: Start 必须早于 End")
            if start < prev_end:
                fail(f"{sheet_name} 第{rn}行: 与上一段时间重叠")
            prev_end = end
            loc = str(r.get("Location") or "").strip()
            if not loc:
                fail(f"{sheet_name} 第{rn}行: Location 为空")
            entries.append({
                "StartMinutes": start,
                "EndMinutes": end,
                "Location": loc,
                "Action": str(r.get("Action") or "").strip(),
                "Danger": int(num(r.get("Danger"), f"{sheet_name} 第{rn}行.Danger", lo=0, hi=3)),
                "Note": str(r.get("Note") or ""),
            })
        schedules.append({"FamilyId": family_id, "Entries": entries})
        if entries and (entries[0]["StartMinutes"] > global_cfg["DayStartMinutes"]):
            warn(f"{sheet_name}: 第一条动线晚于一天开始时刻")

    # ---- Tasks（当前任务，2026-10-14 草案） ----
    headers, rows = read_table(wb["Tasks"], "Tasks")
    tasks = []
    seen_task_ids = set()
    seen_orders = set()
    for r in rows:
        rn = r["_row"]
        task_id = str(r.get("TaskId") or "").strip()
        if not task_id:
            fail(f"Tasks 第{rn}行: TaskId 为空")
            continue
        if task_id in seen_task_ids:
            fail(f"Tasks 第{rn}行: TaskId '{task_id}' 重复")
        seen_task_ids.add(task_id)
        if not (task_id.isascii() and task_id[0].isalpha()
                and all(ch.isalnum() or ch == "_" for ch in task_id)):
            fail(f"Tasks 第{rn}行: TaskId '{task_id}' 必须是英文标识符（字母开头）")

        order = int(num(r.get("Order"), f"Tasks.{task_id}.Order", lo=1, hi=999))
        if order in seen_orders:
            fail(f"Tasks.{task_id}: Order {order} 与其它任务重复")
        seen_orders.add(order)

        title = str(r.get("Title") or "").strip()
        if not title:
            fail(f"Tasks.{task_id}: Title 为空")

        # TriggerItemId 可以留空（无触发道具的任务，靠剧情/调试推进）；填了就必须在道具表里能查到。
        # 2026-10-18：支持逗号分隔多件（含义 = 全部拥有才完成，顺序无关）。
        trigger = str(r.get("TriggerItemId") or "").strip()
        if trigger:
            for one in [x.strip() for x in trigger.split(",") if x.strip()]:
                if not any(it["ItemId"] == one for it in items):
                    fail(f"Tasks.{task_id}.TriggerItemId: 道具表里没有 '{one}'")

        tasks.append({
            "TaskId": task_id,
            "Order": order,
            "Title": title,
            "Detail": str(r.get("Detail") or "").strip(),
            "TriggerItemId": trigger,
            "OnCompleteStory": str(r.get("OnCompleteStory") or "").strip(),
            "Note": str(r.get("Note") or ""),
        })

    # Order 必须从 1 开始连续（任务是一条直线，中间不能断号，否则 UI 推进会出现空洞）
    if tasks:
        tasks.sort(key=lambda t: t["Order"])
        for i, t in enumerate(tasks, start=1):
            if t["Order"] != i:
                fail(f"Tasks: Order 必须从 1 起连续，第 {i} 条却是 {t['Order']}（TaskId={t['TaskId']}）")

    # ---- Story（剧情节点，2026-10-16 开场竖切草案） ----
    # 文本一律来自本表（不硬编码在 C++ 里），策划只改 Excel 即可换文案。
    # Effect 语法（分号分隔多个，顺序执行）：
    #   Mood+5 / Mood-5   躁郁值增减（剧情内一律 ±5，见说明 sheet）
    #   Sus+5  / Sus-5    怀疑值增减
    #   Give:ItemId:N     给道具（N 省略即 1）
    #   Task:TaskId       把当前任务切到该任务
    #   Death             死亡 -> 走现有周目结算（LoopSubsystem::RequestLoopEnd(Caught)）
    #   Win               标记本局通关（2026-10-18）：**剧情结束时**才显示通关界面
    #                     （放在最后一句文本的节点上，玩家读完再收尾）
    #                     ⚠️ 原剧情的"出去"分支（蜘蛛网/棉棒铁丝判定/大雪/强风操控）已按用户要求
    #                     整段删除，改成"从门缝出去 = 直接通关"（场景里 PlayTable 的 Meta.Spot:EscapeHole）
    # Condition 语法：空（无条件）/ Stage>=2 / Stage<=1 / HasItem:ItemId
    # 玩法门（2026-10-16：分支抉择 = 回到游戏里做一件事，而不是文字菜单）：
    #   GateKind       Spot / Distance / Interact / Timeout（空 = 普通文本节点）
    #   GateTarget     Spot/Interact: 场景标签 Meta.Spot:<Id>；Distance: 留空 = 以门激活时玩家位置为参照
    #   GateParam      Spot: 判定外扩 cm；Distance: cm（>0 至少这么远 / <0 靠近到 |N| 内）；Interact: 可交互距离
    #   GateTimeLimit  限时秒（0 = 不限时）
    #   GateNext       满足 -> 该节点；GateTimeoutNext 超时 -> 该节点（留空则用 Next）
    #   ObjectiveText  门激活期间任务 HUD 上显示的目标文字（如「躲到床下」）
    #   GateHoldSeconds  满足后还要"连续保持"多少秒才算完成（0/空 = 一满足就完成）。
    #                    中途条件失效（离开区域等）则计时清零，需重新累计。
    #                    GateTimeLimit 的语义 = **从门激活起算的总时限**，所以必须 > GateHoldSeconds。
    #   ObjectiveHoldText 保持阶段显示的任务 HUD 文案，支持 {N} = 剩余秒数（只支持这一个占位符）；留空则用 ObjectiveText
    headers, rows = read_table(wb["Story"], "Story")
    story = []
    seen_nodes = set()
    valid_effects = ("Mood", "Sus", "Give", "Task", "Death", "Win")
    for r in rows:
        rn = r["_row"]
        node_id = str(r.get("NodeId") or "").strip()
        if not node_id:
            fail(f"Story 第{rn}行: NodeId 为空")
            continue
        if node_id in seen_nodes:
            fail(f"Story 第{rn}行: NodeId '{node_id}' 重复")
        seen_nodes.add(node_id)
        if not (node_id.isascii() and node_id[0].isalpha()
                and all(ch.isalnum() or ch == "_" for ch in node_id)):
            fail(f"Story 第{rn}行: NodeId '{node_id}' 必须是英文标识符（字母开头）")

        text = str(r.get("Text") or "").replace("\\n", "\n")   # 表里写 \n -> JSON 里是真换行
        if not text.strip():
            fail(f"Story.{node_id}: Text 为空")

        effect = str(r.get("Effect") or "").strip()
        for part in [p.strip() for p in effect.split(";") if p.strip()]:
            head = part.split(":")[0].split("+")[0].split("-")[0]
            if head not in valid_effects:
                fail(f"Story.{node_id}.Effect: 无法识别 '{part}'（合法：{'/'.join(valid_effects)}）")
            if head == "Give":
                seg = part.split(":")
                if len(seg) < 2 or not seg[1].strip():
                    fail(f"Story.{node_id}.Effect: '{part}' 缺少 ItemId（写法 Give:ItemId[:N]）")
                elif not any(it["ItemId"] == seg[1].strip() for it in items):
                    fail(f"Story.{node_id}.Effect: 道具表里没有 '{seg[1].strip()}'")
            if head == "Task":
                seg = part.split(":")
                if len(seg) < 2 or not any(t["TaskId"] == seg[1].strip() for t in tasks):
                    fail(f"Story.{node_id}.Effect: 任务表里没有 '{part.split(':')[-1].strip()}'")

        cond = str(r.get("Condition") or "").strip()
        if cond and not (cond.startswith("Stage") or cond.startswith("HasItem:") or cond.startswith("NoItem:")):
            fail(f"Story.{node_id}.Condition: 无法识别 '{cond}'（合法：Stage>=N / Stage<=N / HasItem:ItemId / NoItem:ItemId）")
        if cond.startswith("NoItem:"):
            cond_item = cond.split(":", 1)[1].strip()
            if not any(it["ItemId"] == cond_item for it in items):
                fail(f"Story.{node_id}.Condition: 道具表里没有 '{cond_item}'")

        a_text = str(r.get("ChoiceA_Text") or "").strip()
        b_text = str(r.get("ChoiceB_Text") or "").strip()
        a_next = str(r.get("ChoiceA_Next") or "").strip()
        b_next = str(r.get("ChoiceB_Next") or "").strip()
        if bool(a_text) != bool(b_text):
            fail(f"Story.{node_id}: 有 ChoiceA 就必须有 ChoiceB（两选项成对出现）")
        if a_text and not (a_next and b_next):
            fail(f"Story.{node_id}: 有选项就必须填 ChoiceA_Next 和 ChoiceB_Next")
        if bool(a_next) != bool(b_next):
            fail(f"Story.{node_id}: ChoiceA_Next / ChoiceB_Next 必须成对填写")

        # ---- 玩法门 ----
        gate_kind = str(r.get("GateKind") or "").strip()
        gate_target = str(r.get("GateTarget") or "").strip()
        gate_next = str(r.get("GateNext") or "").strip()
        gate_timeout_next = str(r.get("GateTimeoutNext") or "").strip()
        objective = str(r.get("ObjectiveText") or "").strip()
        # UnderMesh（2026-10-17 新增）= 真的"藏在目标物件下方"：从玩家向上打一条竖直线，
        # 只有命中目标物件才算（Spot 的包围盒会把床周围一大片都算进去，实测连重生点都在盒边）。
        # Region（2026-10-17 新增）= 显式 AABB 判定（用户画了矩形框，要求"退到框的右半边"）：
        # GateRegionBox 写 "XMin,XMax,YMin,YMax,ZMin,ZMax"，世界坐标 cm。
        valid_gates = ("Spot", "Region", "UnderMesh", "Distance", "Interact", "Timeout")
        if gate_kind and gate_kind not in valid_gates:
            fail(f"Story.{node_id}.GateKind: '{gate_kind}' 不在 {valid_gates} 内")
        region_box_text = str(r.get("GateRegionBox") or "").strip()
        if gate_kind:
            if not gate_next:
                fail(f"Story.{node_id}: 玩法门必须填 GateNext（满足后去哪）")
            if gate_kind in ("Region", "Spot", "UnderMesh", "Interact") and not gate_target:
                if gate_kind == "Region":
                    pass  # Region 的 GateTarget 可为空：为空时就必须给 GateRegionBox（下面会校验）
                else:
                    fail(f"Story.{node_id}: GateKind={gate_kind} 必须填 GateTarget（区域 Id；有标签时优先 Meta.Spot:<Id>，否则按 SpotMeshNames 的网格名）")
            # Region 有两种给区域的方式（2026-10-17）：
            #   ① GateTarget 指向场景里打了标签 Meta.Spot:<Id> 的 Actor（区域=它的运行时包围盒，地编挪动 Actor 即可改区域）
            #   ② GateRegionBox 直接写死 AABB
            # 两者至少有一个；两个都给时运行期是"标签优先、配表兜底"。
            if gate_kind == "Region" and not region_box_text and not gate_target:
                fail(f"Story.{node_id}: GateKind=Region 需要 GateTarget（场景标签 Meta.Spot:<Id>）或 GateRegionBox（\"XMin,XMax,YMin,YMax,ZMin,ZMax\"）二者之一")
            if gate_kind == "Region" and not region_box_text and gate_target:
                warn(f"Story.{node_id}: Region 门没写 GateRegionBox，依赖场景里打了标签 Meta.Spot:{gate_target} 的 Actor 的包围盒"
                     "（该 Actor 不存在时这道门永远不满足、只能等超时）")
            if not objective:
                fail(f"Story.{node_id}: 玩法门必须填 ObjectiveText（任务 HUD 上给玩家看的指引）")
            if a_text:
                fail(f"Story.{node_id}: 同一节点不要同时用「选项按钮」和「玩法门」（会同时出现两套交互）")
            if not str(r.get("Next") or "").strip() and not gate_timeout_next:
                warn(f"Story.{node_id}: 玩法门没有 Next 也没有 GateTimeoutNext，超时后剧情会结束")
        elif gate_target or gate_next or gate_timeout_next or objective \
                or region_box_text or str(r.get("GateHoldSeconds") or "").strip() \
                or str(r.get("ObjectiveHoldText") or "").strip():
            fail(f"Story.{node_id}: 填了 Gate* / ObjectiveText 却没填 GateKind")

        # GateRegionBox 校验：必须是 6 个数字、每个数字可解析、且 min < max
        if region_box_text:
            segs = [s.strip() for s in region_box_text.split(",")]
            if len(segs) != 6:
                fail(f"Story.{node_id}.GateRegionBox: 必须是 6 个数字 'XMin,XMax,YMin,YMax,ZMin,ZMax'（当前 {len(segs)} 段）")
            vals = []
            for s in segs:
                try:
                    vals.append(float(s))
                except ValueError:
                    fail(f"Story.{node_id}.GateRegionBox: '{s}' 不是数字")
            if vals[0] > vals[1] or vals[2] > vals[3] or vals[4] > vals[5]:
                fail(f"Story.{node_id}.GateRegionBox: XMin>XMax / YMin>YMax / ZMin>ZMax 写反了（当前 '{region_box_text}'）")
            if gate_kind and gate_kind != "Region":
                warn(f"Story.{node_id}: 填了 GateRegionBox 但 GateKind={gate_kind}（只有 Region 门会用到它）")

        gate_param = num(r.get("GateParam"), f"Story.{node_id}.GateParam", lo=-100000, hi=100000) if gate_kind else 0.0
        gate_limit = num(r.get("GateTimeLimit"), f"Story.{node_id}.GateTimeLimit", lo=0, hi=3600) if gate_kind else 0.0
        gate_hold = 0.0
        raw_hold = r.get("GateHoldSeconds")
        if gate_kind and raw_hold is not None and str(raw_hold).strip() != "":
            gate_hold = num(raw_hold, f"Story.{node_id}.GateHoldSeconds", lo=0, hi=600)
        hold_text = str(r.get("ObjectiveHoldText") or "").strip() if gate_kind else ""
        # 保持阶段只能用于"条件型"门（Timeout 门本身就是等时间，没有"保持"概念）
        if gate_hold > 0 and gate_kind not in ("Spot", "Region", "UnderMesh", "Distance", "Interact"):
            fail(f"Story.{node_id}: GateHoldSeconds 只能用于 Spot/Region/UnderMesh/Distance/Interact 门（{gate_kind} 不支持保持）")
        # 总时限必须大于保持时长，否则这道门永远无法完成
        if gate_hold > 0 and gate_limit > 0 and gate_hold >= gate_limit:
            fail(f"Story.{node_id}: GateHoldSeconds({gate_hold:g}) 必须小于 GateTimeLimit({gate_limit:g})，否则门永远完不成")
        # ObjectiveHoldText 只允许 {N} 占位符
        if hold_text:
            probe = hold_text.replace("{N}", "")
            if "{" in probe or "}" in probe:
                fail(f"Story.{node_id}.ObjectiveHoldText: 只支持 {{N}} 占位符（当前 '{hold_text}'）")
            if gate_hold <= 0:
                warn(f"Story.{node_id}: 填了 ObjectiveHoldText 但 GateHoldSeconds=0（保持文案不会用到）")

        story.append({
            "NodeId": node_id,
            "Scene": str(r.get("Scene") or "").strip(),
            "Speaker": str(r.get("Speaker") or "").strip(),
            "Text": text,
            "Next": str(r.get("Next") or "").strip(),
            "ChoiceA_Text": a_text,
            "ChoiceA_Next": a_next,
            "ChoiceB_Text": b_text,
            "ChoiceB_Next": b_next,
            "Effect": effect,
            "Condition": cond,
            "GateKind": gate_kind,
            "GateRegionBox": region_box_text,
            "GateTarget": gate_target,
            "GateParam": gate_param,
            "GateTimeLimit": gate_limit,
            "GateNext": gate_next,
            "GateTimeoutNext": gate_timeout_next,
            "ObjectiveText": objective,
            "GateHoldSeconds": gate_hold,
            "ObjectiveHoldText": hold_text,
            "Note": str(r.get("Note") or ""),
        })

    # 交叉校验：所有 Next / ChoiceNext / GateNext 必须指向存在的节点（或留空 = 结束）
    node_ids = {n["NodeId"] for n in story}
    for n in story:
        for field in ("Next", "ChoiceA_Next", "ChoiceB_Next", "GateNext", "GateTimeoutNext"):
            target = n[field]
            if target and target not in node_ids:
                fail(f"Story.{n['NodeId']}.{field}: 指向不存在的节点 '{target}'")

    # 交叉校验：任务的 OnCompleteStory 必须指向存在的剧情节点（任务表在 Story 之前解析，所以放在这里）
    for t in tasks:
        target = t.get("OnCompleteStory") or ""
        if target and target not in node_ids:
            fail(f"Tasks.{t['TaskId']}.OnCompleteStory: 剧情表里没有节点 '{target}'")

    # 交叉校验：道具的 FollowUpStoryNode（剧情道具读完后播的台词节点）必须存在，
    # 且只对剧情道具（Category=Story）有意义 —— 别的道具填了也不会走到。
    for it in items:
        target = it.get("FollowUpStoryNode") or ""
        if not target:
            continue
        if target not in node_ids:
            fail(f"Items.{it['ItemId']}.FollowUpStoryNode: 剧情表里没有节点 '{target}'")
        if it["Category"] != "Story":
            warn(f"Items.{it['ItemId']}.FollowUpStoryNode: 该道具 Category={it['Category']}（不是 Story），"
                 "这条台词不会播（剧情道具才会先播回忆文本）")

    # ---- 汇总 ----
    cfg = {"Global": global_cfg, "Suspicion": sus, "Mood": mood, "Periods": periods, "Schedules": schedules,
           "Items": items, "Tasks": tasks, "Story": story}

    if errors:
        print("校验失败：")
        for e in errors:
            print("  [ERROR]", e)
        sys.exit(1)
    for w in warnings:
        print("  [WARN]", w)

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "game_config.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    print("exported:", out)

    # ---- 同时生成"编译进二进制"的内嵌版本 ----
    # 为什么需要：打 Android APK 时 Content/Config/game_config.json 属于"非资产文件"，
    # DirectoriesToAlwaysStageAsNonUFS / AsUFS 两种打包设置都没能把它带进 APK
    # （重打后 APK 条目数完全不变），真机上任务栏/交互全空。
    # 所以把同一份 JSON 作为 C++ 窄字符原始字符串字面量生成到 .inl，由 Config 子系统在
    # "文件不可用"时回退解析 -> 一定带得上。
    # 注意：改完配表必须【重新编译】一次，内嵌内容才会更新（编辑器里仍优先读文件，热改不受影响）。
    embed_config_json(out, cfg)


# 内嵌用的原始字符串定界符：足够特别，生成前校验 JSON 里不含它（否则会截断字面量）
EMBED_DELIM = "META_CFG_V1"
EMBED_REL = os.path.join("Source", "Metamorphosis", "Config", "GameConfigJson.inl")


def embed_config_json(json_path, cfg):
    """把刚导出的 JSON 再写成 GameConfigJson.inl（供 C++ 在文件不可用时回退解析）。"""
    # 与落盘完全一致的格式，方便运行时做"内嵌 vs 文件"一致性对比
    text = json.dumps(cfg, ensure_ascii=False, indent=2)
    if (')' + EMBED_DELIM + '"') in text:
        print("  [ERROR] 配表内容里出现了内嵌定界符 )%s\"，会截断字符串字面量，请修改内容后再导出" % EMBED_DELIM)
        sys.exit(1)
    text = text.replace("\r", "")

    inl_path = os.path.join(ROOT, EMBED_REL)
    os.makedirs(os.path.dirname(inl_path), exist_ok=True)

    # MSVC 对单个字面量（以及相邻字面量拼接后的总长）都有上限（error C2026: string too big）。
    # 因此：切成很多小块（每块约 2000 字节），每块一个独立的 C 字符串，代码里运行时拼起来。
    # 关键：必须按【UTF-8 字符边界】切，否则某块会是非法 UTF-8，编译器/源文件都会出问题。
    def split_utf8(s, max_bytes):
        out, cur, n = [], [], 0
        for ch in s:
            b = len(ch.encode("utf-8"))
            if n + b > max_bytes and cur:
                out.append("".join(cur))
                cur, n = [], 0
            cur.append(ch)
            n += b
        if cur:
            out.append("".join(cur))
        return out or [""]

    chunks = split_utf8(text, 2000)
    data = text.encode("utf-8")

    with open(inl_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("// 本文件由 Tools/export_config.py 自动生成，请勿手改。\n")
        f.write("// 作用：把 Content/Config/game_config.json 同一份内容编译进二进制，供\n")
        f.write("//       UMetamorphosisConfigSubsystem 在\"文件不可用\"（典型：Android APK 里非资产文件没被带上）时回退解析。\n")
        f.write("// 重新生成：改 ConfigSource/数值配置表.xlsx -> python Tools/export_config.py -> 重新编译。\n\n")
        f.write("#pragma once\n\n")
        f.write("// UTF-8 窄字符原始字符串，分 %d 块、运行时拼接（MSVC 对字面量总长有上限，不能写成一个大字符串）。\n"
                % len(chunks))
        f.write("// 不用 TEXT()：TEXT 是 L 拼接宏，和原始字符串字面量组合在部分编译器上不可靠。\n")
        f.write("static const char* const GEmbeddedGameConfigJsonChunks[] =\n{\n")
        for i, c in enumerate(chunks):
            delim = "%s_%03d" % (EMBED_DELIM, i)
            if (')' + delim + '"') in c:
                print("  [ERROR] 配表第 %d 块里出现了定界符，请调整内容" % i)
                sys.exit(1)
            f.write('\tR"%s(%s)%s",\n' % (delim, c, delim))
        f.write("};\n")
        f.write("static const int32 GEmbeddedGameConfigJsonChunkCount = %d;\n" % len(chunks))
        f.write("static const int32 GEmbeddedGameConfigJsonUtf8Len = %d; // 拼接后的字节数（不含结尾 0）\n"
                % len(data))
    print("exported(embedded):", inl_path, "(%d bytes json, %d chunks)" % (len(data), len(chunks)))

if __name__ == "__main__":
    main()
